"""Read-only Excel metadata audit; never exports or changes a source workbook.

Run: python -X utf8 tools/audit-import-sources.py
The JSON contains headers, types and counts, not personal source records.
"""
import json
import re
import warnings
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


def profile(workbook, sheet, header_row, columns, first_row, last_row, key_col, key_pattern):
    ws = workbook[sheet]
    indexes = {col: column_index_from_string(col) for col in columns}
    max_col = max(max(indexes.values()), key_col)
    rows = list(ws.iter_rows(min_row=header_row, max_row=last_row, max_col=max_col))
    headers = rows[0]
    records = [row for row in rows[first_row - header_row:]
               if re.fullmatch(key_pattern, str(row[key_col - 1].value or "").strip())]
    fields = []
    for col, index in indexes.items():
        cells = [row[index - 1] for row in records]
        nonblank = [cell for cell in cells if cell.value is not None]
        fields.append({
            "column": col,
            "header_ref": f"{sheet}!{col}{header_row}",
            "header": headers[index - 1].value,
            "excel_types": dict(Counter(cell.data_type for cell in nonblank)),
            "number_formats": dict(Counter(cell.number_format for cell in nonblank)),
            "blank": len(cells) - len(nonblank),
        })
    return {"sheet": sheet, "scan_rows": [first_row, last_row],
            "record_filter": {"column": key_col, "pattern": key_pattern},
            "records": len(records), "fields": fields}


def main():
    inventory = []
    for path in sorted(DOCS.rglob("*.xlsx")):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        inventory.append({"file": path.relative_to(ROOT).as_posix(), "sheets": wb.sheetnames})
        wb.close()

    selections = [
        ("Thông tin nhà.xlsx", [
            ("NVVH", 1, list("BCDEFGHIJKLMNOPQRS"), 3, 205, 2, r"[TSG]\d+\w*"),
            ("Địa chỉ chi tiết", 1, list("ABCDE"), 2, 205, 1, r"[TSG]\d+\w*"),
        ]),
        ("Danh sách mã HĐ điện nước mạng.xlsx", [
            ("năm 2026", 3, ["A", "B", "F", "O", "P", "Y", "Z"], 4, 205, 1, r"[TSG]\d+\w*"),
        ]),
        ("Hoá đơn tiền nhà tháng 9 và dịch vụ tháng 9 năm 2026.xlsx", [
            (sheet, 2, ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "AY", "BA", "BB", "BC"],
             5, 2000, 2, r"[TSG]\d+\w*") for sheet in ["NHÀ T", "NHÀ G", "NHÀ S"]
        ]),
        ("Hoa hồng năm 2025-2026 (1).xlsx", [
            ("HOA HỒNG THÁNG 9.26", 2, list("DEFGHNO"), 4, 228, 4, r"\d+[TSG]\d+\w*"),
        ]),
        ("G1.31.8.26.xlsx", [
            ("HĐ T8.26  ", 1, list("ABCDEFGHIJKLMNO"), 4, 19, 2, r"G\d+\w*"),
        ]),
    ]
    profiles = []
    for filename, specs in selections:
        wb = openpyxl.load_workbook(DOCS / filename, read_only=True, data_only=False)
        profiles.append({"file": f"docs/{filename}", "profiles": [profile(wb, *spec) for spec in specs]})
        wb.close()
    output = DOCS / "qa" / "import-schema" / "source-metadata.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps({
        "date": "2026-10-21", "base_commit": "70763e39fbd335b47c53fea972fdaf3522aad63e",
        "mode": "read-only", "type_key": {"s": "text", "n": "number", "d": "date", "f": "formula", "b": "boolean", "e": "Excel error"},
        "limits": "Literal cell types and number formats only. Formula caches are not evaluated. Bounded row scans and building/room filters are recorded per sheet. Inventory is not a field-level audit of every sheet.",
        "inventory": inventory, "sources": profiles,
    }, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(f"Inventoried {len(inventory)} workbooks; profiled {sum(len(p['profiles']) for p in profiles)} sheets in {len(profiles)} sources.")
    print(output.relative_to(ROOT).as_posix())


if __name__ == "__main__":
    main()
